"""Small helpers on top of python-docx used by build_report.py.

Keeps the report script readable: headings, paragraphs with **bold** / `code`
markup, bullet lists, numbered captions, tables, code and transcript blocks,
and a Table of Contents that is a real Word TOC field pre-filled with entries
and page numbers (measured from the rendered PDF in a second pass).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x2F, 0x5F)
BLUE = RGBColor(0x5B, 0x84, 0xC4)
GREY = RGBColor(0x59, 0x59, 0x59)
BODY_FONT = "Times New Roman"
SANS = "Arial"
MONO = "Courier New"
# **bold**, `code`, *italic* (italic must not start/end with a space, so "a * b" is left alone)
MARKUP = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*[^*\s]\*)")


def _shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _set_cell_bg(cell, fill: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), fill)


def _font(run, name=BODY_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def add_field(paragraph, instr: str, placeholder: str = "1"):
    """Insert a simple Word field (e.g. PAGE) into a paragraph."""
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        run = paragraph.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
            run._r.append(el)
        elif kind == "text":
            run.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        _font(run, SANS, 8.5, color=GREY)
    return paragraph


@dataclass
class Report:
    doc: Document = field(default_factory=Document)
    fig_no: dict = field(default_factory=dict)
    tab_no: dict = field(default_factory=dict)
    chapter: int = 0
    section: int = 0
    toc: list = field(default_factory=list)        # (level, text)
    figures: list = field(default_factory=list)    # caption text
    tables: list = field(default_factory=list)

    # ------------------------------------------------------------ setup
    def setup(self, footer_text: str) -> None:
        d = self.doc
        sec = d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin = sec.bottom_margin = Cm(2.3)
        sec.left_margin = sec.right_margin = Cm(2.4)
        sec.footer_distance = Cm(1.1)
        st = d.styles["Normal"]
        st.font.name, st.font.size = BODY_FONT, Pt(11.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        pf = st.paragraph_format
        pf.space_after, pf.line_spacing = Pt(6), 1.15
        for lvl, size, before in ((1, 17, 0), (2, 13.5, 12), (3, 12, 9)):
            hs = d.styles[f"Heading {lvl}"]
            hs.font.name, hs.font.size, hs.font.bold = BODY_FONT, Pt(size), True
            hs.font.color.rgb = NAVY
            rf = hs.element.rPr.rFonts
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(attr)) is not None:
                    del rf.attrib[qn(attr)]
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(attr), BODY_FONT)
            hs.paragraph_format.space_before = Pt(before)
            hs.paragraph_format.space_after = Pt(6 if lvl > 1 else 12)
            hs.paragraph_format.keep_with_next = True
        d.styles["Heading 1"].paragraph_format.page_break_before = True
        for name in ("List Bullet", "List Number"):
            d.styles[name].font.name = BODY_FONT
            d.styles[name].paragraph_format.space_after = Pt(3)
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(f"{footer_text} | Page "), SANS, 8.5, color=GREY)
        add_field(p, "PAGE")

    # ------------------------------------------------------------ text
    def para(self, text: str = "", align="justify", size=None, italic=False, space_after=None, keep_next=False):
        p = self.doc.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = keep_next
        self.inline(p, text, size=size, italic=italic)
        return p

    def inline(self, p, text: str, size=None, italic=False):
        """Render **bold**, *italic* and `code` markup."""
        for part in MARKUP.split(text):
            if not part:
                continue
            if part.startswith("**"):
                _font(p.add_run(part[2:-2]), size=size, bold=True, italic=italic)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                _font(p.add_run(part[1:-1]), size=size, italic=True)
            elif part.startswith("`"):
                _font(p.add_run(part[1:-1]), MONO, (size or 11.5) - 1.5)
            else:
                _font(p.add_run(part), size=size, italic=italic)

    def bullets(self, items: list[str], numbered=False):
        for n, it in enumerate(items, 1):
            if numbered:
                # explicit numbers: Word's shared "List Number" would continue across sections
                p = self.doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent, pf.first_line_indent, pf.space_after = Cm(0.75), Cm(-0.75), Pt(3)
                pf.tab_stops.add_tab_stop(Cm(0.75))
                _font(p.add_run(f"{n}.\t"))
            else:
                p = self.doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.inline(p, it)

    def h1(self, title: str, numbered=True, label: str | None = None, new_page=True):
        """numbered chapters get 1, 2, ...; appendices pass label='A', 'B', ..."""
        if numbered:
            self.chapter += 1
            self.section = 0
            self.label = str(self.chapter)
            text = f"{self.chapter}  {title}"
        else:
            self.label = label or ""
            text = title
        h = self.doc.add_heading(text, level=1)
        if not new_page:
            h.paragraph_format.page_break_before = False
            h.paragraph_format.space_before = Pt(24)
        self.toc.append((1, text))

    def h2(self, title: str):
        self.section += 1
        self.sub = 0
        text = f"{self.chapter}.{self.section}  {title}"
        self.doc.add_heading(text, level=2)
        self.toc.append((2, text))

    def h3(self, title: str):
        self.sub = getattr(self, "sub", 0) + 1
        self.doc.add_heading(f"{self.chapter}.{self.section}.{self.sub}  {title}", level=3)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ figures / tables
    def _caption(self, kind: str, text: str) -> str:
        counter = self.fig_no if kind == "Figure" else self.tab_no
        key = getattr(self, "label", str(self.chapter))
        counter[key] = counter.get(key, 0) + 1
        label = f"{kind} {key}.{counter[key]}: {text}"
        (self.figures if kind == "Figure" else self.tables).append(label)
        return label

    def figure(self, path: str, caption: str, width_cm: float = 16.0):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width_cm))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(c.add_run(self._caption("Figure", caption)), size=10.5, italic=True, color=GREY)

    def table(self, header: list[str], rows: list[list[str]], caption: str, widths_cm: list[float] | None = None,
              size: float = 9.5):
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_before = Pt(6)
        _font(c.add_run(self._caption("Table", caption)), size=10.5, italic=True, color=GREY)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            _font(cell.paragraphs[0].add_run(h), SANS, size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _set_cell_bg(cell, "1F2F5F")
        # repeat header row on each page
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for part in MARKUP.split(str(val)):
                    if not part:
                        continue
                    if part.startswith("**"):
                        _font(para.add_run(part[2:-2]), size=size, bold=True)
                    elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                        _font(para.add_run(part[1:-1]), size=size, italic=True)
                    elif part.startswith("`"):
                        _font(para.add_run(part[1:-1]), MONO, size - 1)
                    else:
                        _font(para.add_run(part), size=size)
                if r_i % 2 == 1:
                    _set_cell_bg(cells[i], "EEF2F8")
        if not widths_cm:
            widths_cm = [16.2 / len(header)] * len(header)
        t.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(layout)
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w * 567)))
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def code(self, text: str, title: str | None = None, size: float = 8.0):
        """Monospaced block with light shading (used for code and transcripts)."""
        if title:
            p = self.doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(2)
            _font(p.add_run(title), SANS, 9, bold=True, color=NAVY)
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.2)
            if i < len(lines) - 1 and len(lines) <= 40:
                pf.keep_with_next = True
            _shade(p._p.get_or_add_pPr(), "F3F3F3")
            _font(p.add_run(line if line else " "), MONO, size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------------ TOC
    def toc_block(self, title: str, entries: list[tuple[int, str]], pages: dict[str, int] | None, field_instr=None):
        h = self.doc.add_paragraph()
        _font(h.add_run(title), BODY_FONT, 17, bold=True, color=NAVY)
        h.paragraph_format.space_after = Pt(10)
        first = True
        paras = []
        for level, text in entries:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(3 if level == 1 else 1)
            pf.left_indent = Cm(0 if level == 1 else 0.8)
            pf.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if first and field_instr:
                for kind, val in (("begin", None), ("instr", field_instr), ("separate", None)):
                    run = p.add_run()
                    if kind == "instr":
                        el = OxmlElement("w:instrText")
                        el.set(qn("xml:space"), "preserve")
                        el.text = f" {val} "
                    else:
                        el = OxmlElement("w:fldChar")
                        el.set(qn("w:fldCharType"), kind)
                    run._r.append(el)
            first = False
            page = (pages or {}).get(text, "")
            _font(p.add_run(f"{text}\t{page if page != '' else '00'}"), size=11 if level == 1 else 10.5,
                  bold=(level == 1))
            paras.append(p)
        if field_instr and paras:
            run = paras[-1].add_run()
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), "end")
            run._r.append(el)
